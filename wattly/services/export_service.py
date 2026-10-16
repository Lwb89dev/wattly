"""CSV / XLSX export of bills, analytics and comparison results."""

from __future__ import annotations

import csv
import io
from dataclasses import asdict

from wattly.core.analytics import AnalyticsReport
from wattly.core.comparisons import ComparisonResult
from wattly.core.models import Bill

BILL_COLUMNS = [
    ("id", "ID"), ("provider_name", "Provider"), ("account_number", "Account"),
    ("pod_identifier", "POD"), ("issue_date", "Issue date"),
    ("billing_period_start", "Period start"), ("billing_period_end", "Period end"),
    ("total_amount_cents", "Total (cents)"), ("currency", "Currency"),
    ("total_kwh", "kWh"), ("avg_daily_kwh", "Avg daily kWh"),
    ("avg_cost_per_kwh", "Energy €/kWh"),
    ("effective_total_cost_per_kwh", "Effective €/kWh"),
    ("fixed_charges_cents", "Fixed (cents)"), ("energy_charges_cents", "Energy (cents)"),
    ("taxes_and_fees_cents", "Taxes (cents)"),
    ("previous_reading", "Prev reading"), ("current_reading", "Curr reading"),
    ("tariff_type", "Tariff"), ("parse_method", "Parse method"),
    ("parse_confidence", "Confidence"), ("notes", "Notes"),
]


def _fmt(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def bills_csv(bills: list[Bill]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([label for _, label in BILL_COLUMNS])
    for b in bills:
        data = asdict(b)
        writer.writerow([_fmt(data.get(k)) for k, _ in BILL_COLUMNS])
    return buf.getvalue()


def monthly_csv(report: AnalyticsReport) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Period", "kWh", "Cost (cents)", "Avg daily kWh", "Effective €/kWh",
                     "Rolling avg kWh (3)"])
    for i, p in enumerate(report.series):
        writer.writerow([
            p.period_start.isoformat(), f"{p.kwh:.1f}", p.amount_cents,
            f"{p.daily_kwh:.3f}", f"{p.cost_per_kwh:.4f}",
            report.rolling_avg_kwh[i] if i < len(report.rolling_avg_kwh) else "",
        ])
    if report.band_labels:
        writer.writerow([])
        writer.writerow(["Band", "Total kWh", "Total (cents)"])
        for lab in report.band_labels:
            writer.writerow([lab, report.band_kwh.get(lab, 0.0), report.band_cents.get(lab, 0)])
    return buf.getvalue()


def comparison_csv(res: ComparisonResult) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Metric", res.label_a, res.label_b, "Delta (B−A)", "Delta %"])
    rows = [
        ("Total kWh", f"{res.a.kwh:.1f}", f"{res.b.kwh:.1f}", f"{res.kwh_delta:.1f}", res.kwh_delta_pct),
        ("Total cost (cents)", res.a.cents, res.b.cents, res.cost_delta_cents, res.cost_delta_pct),
        ("Avg daily kWh", f"{res.a.daily_kwh:.3f}", f"{res.b.daily_kwh:.3f}",
         f"{res.daily_kwh_delta:.3f}", res.daily_kwh_delta_pct),
        ("Effective €/kWh",
         f"{res.a.cost_per_kwh:.4f}" if res.a.cost_per_kwh else "",
         f"{res.b.cost_per_kwh:.4f}" if res.b.cost_per_kwh else "",
         f"{res.cost_per_kwh_delta:.4f}" if res.cost_per_kwh_delta is not None else "", ""),
    ]
    writer.writerows(rows)
    writer.writerow([])
    writer.writerow([f"Bills in {res.label_a}"] + [f"{b.issue_date}: {b.total_kwh} kWh" for b in res.a.bills])
    writer.writerow([f"Bills in {res.label_b}"] + [f"{b.issue_date}: {b.total_kwh} kWh" for b in res.b.bills])
    return buf.getvalue()


def workbook_xlsx(
    bills: list[Bill],
    report: AnalyticsReport | None = None,
    comparison: ComparisonResult | None = None,
) -> bytes:
    """Multi-sheet Excel workbook: Bills (+bands), Monthly, Bands, Comparison."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append([label for _, label in BILL_COLUMNS] + [f"{lab} kWh" for lab in (report.band_labels if report else [])])
    for b in bills:
        data = asdict(b)
        row = [_fmt(data.get(k)) for k, _ in BILL_COLUMNS]
        if report:
            by_label = {band.label: band.kwh for band in b.time_bands}
            row += [by_label.get(lab) for lab in report.band_labels]
        ws.append(row)
    if report is not None:
        ws2 = wb.create_sheet("Monthly")
        ws2.append(["Period", "kWh", "Cost (cents)", "Avg daily kWh", "Effective €/kWh"])
        for p in report.series:
            ws2.append([p.period_start.isoformat(), round(p.kwh, 1), p.amount_cents,
                        round(p.daily_kwh, 3), round(p.cost_per_kwh, 4)])
        if report.band_labels:
            ws3 = wb.create_sheet("Bands")
            ws3.append(["Band", "Total kWh", "Total (cents)"])
            for lab in report.band_labels:
                ws3.append([lab, report.band_kwh.get(lab, 0.0), report.band_cents.get(lab, 0)])
    if comparison is not None:
        ws4 = wb.create_sheet("Comparison")
        ws4.append(["Metric", comparison.label_a, comparison.label_b, "Delta (B−A)", "Delta %"])
        ws4.append(["Total kWh", round(comparison.a.kwh, 1), round(comparison.b.kwh, 1),
                    comparison.kwh_delta, comparison.kwh_delta_pct])
        ws4.append(["Total cost (cents)", comparison.a.cents, comparison.b.cents,
                    comparison.cost_delta_cents, comparison.cost_delta_pct])
        ws4.append(["Avg daily kWh", round(comparison.a.daily_kwh, 3),
                    round(comparison.b.daily_kwh, 3),
                    round(comparison.daily_kwh_delta, 3), comparison.daily_kwh_delta_pct])
        ws4.append(["Effective €/kWh", comparison.a.cost_per_kwh, comparison.b.cost_per_kwh,
                    comparison.cost_per_kwh_delta, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
